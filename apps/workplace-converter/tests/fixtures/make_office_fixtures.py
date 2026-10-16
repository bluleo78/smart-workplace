"""합성 픽스처 sample.doc·sample.xls 를 만드는 1회성 스크립트(테스트가 아님 — 결과 바이너리를 커밋한다).

왜 스크립트로 남기나: .doc/.xls 는 python 라이브러리로 직접 쓸 수 없어 python-docx/openpyxl 로 docx/xlsx 를 만든 뒤
컨테이너 안 soffice 로 구형 형식으로 바꾼다. 출처(합성)를 재현할 수 있게 절차를 코드로 남긴다.

사용(converter 이미지 안, 출력 디렉터리를 마운트):
    soffice·python-docx·openpyxl 이 있는 환경에서 `python make_office_fixtures.py <출력 디렉터리>`
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook


def _soffice_convert(src: Path, fmt: str, out_dir: Path) -> None:
    """soffice 1회 실행으로 src 를 fmt(예 'doc:MS Word 97') 로 바꿔 out_dir 에 쓴다."""
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{profile}",
                        "--convert-to", fmt, "--outdir", str(out_dir), str(src)], check=True, timeout=120)


def main(out_dir: Path) -> None:
    """한글·영문 혼합 문단을 담은 sample.doc 와 A1=`변환 테스트` 인 sample.xls 를 만든다."""
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)

        doc = Document()
        doc.add_heading("변환 테스트 문서 (Conversion test)", level=1)
        doc.add_paragraph("이 문서는 workplace-converter 엔진 테스트용 합성 문서입니다.")
        doc.add_paragraph("AI 시대, 참가자 2,603명이 English 와 한국어를 섞어 씁니다.")
        doc.add_paragraph("The quick brown fox jumps over the lazy dog.")
        docx = work / "sample.docx"
        doc.save(docx)
        _soffice_convert(docx, "doc:MS Word 97", out_dir)

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "변환 테스트"
        ws["B1"] = "Conversion test"
        ws["A2"] = 2603
        ws["B2"] = "명"
        xlsx = work / "sample.xlsx"
        wb.save(xlsx)
        _soffice_convert(xlsx, "xls:MS Excel 97", out_dir)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
