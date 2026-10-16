"""상주 LibreOffice(Office) 관리 테스트 — 동시 1건 슬롯·헬퍼 호출·인프라 실패 시 내리기·N건마다 재기동·공통 마감.

단위 테스트는 실제 soffice 없이 Office.start/stop 과 proc.run 을 가짜로 바꿔 상태 전이만 검증한다.
`@pytest.mark.engine` 테스트는 soffice·python3-uno 가 있는 컨테이너 test 타깃에서만 돈다.
"""
import io
import socket
import subprocess
import time
from pathlib import Path

import pymupdf
import pytest

from app import config
from app.engines import proc, soffice
from app.errors import ConvertError


# 가짜 '떠 있는 soffice' — Office._running 은 _proc 유무로 정해지므로 아무 객체나 채워 기동 상태를 흉내 낸다
RUNNING = object()


@pytest.fixture
def office(monkeypatch, tmp_path):
    """start/stop 을 기록만 하는 가짜로 바꾼 새 Office 인스턴스와 호출 기록."""
    monkeypatch.setenv("TMPDIR", str(tmp_path))
    o = soffice.Office()
    events: list[str] = []

    def fake_start(wait=soffice.START_WAIT_SEC):
        events.append("start")
        o._proc = RUNNING

    def fake_stop():
        events.append("stop")
        o._proc = None

    monkeypatch.setattr(o, "start", fake_start)
    monkeypatch.setattr(o, "stop", fake_stop)
    # 생존 판정은 _running(_proc 유무)만 본다(실제 프로세스·포트 없음)
    monkeypatch.setattr(o, "_alive", lambda: o._running)
    # 이미 떠 있는 상태로 시작 — 첫 변환에서 지연 기동이 끼지 않게
    o._proc = RUNNING
    return o, events


@pytest.fixture
def missing_binary(monkeypatch, tmp_path):
    """soffice 실행 파일이 없는 환경의 실제 Office(start/stop 을 가짜로 바꾸지 않음)."""
    monkeypatch.setenv("TMPDIR", str(tmp_path))
    monkeypatch.setattr(config, "SOFFICE_BIN", str(tmp_path / "no-such-soffice"))
    return soffice.Office()


@pytest.fixture
def helper_ok(monkeypatch):
    """성공하는 가짜 헬퍼 — --out 경로에 결과를 쓰고 호출 명령을 기록한다."""
    calls: list[list[str]] = []

    def run(cmd, timeout, cwd, on_fail=None):
        calls.append(list(cmd))
        Path(cmd[cmd.index("--out") + 1]).write_bytes(b"%PDF-fake")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    monkeypatch.setattr(proc, "run", run)
    return calls


def test_busy_returns_429_without_waiting(office, helper_ok):
    o, _ = office
    assert o._lock.acquire(blocking=False)
    try:
        t0 = time.monotonic()
        with pytest.raises(ConvertError) as e:
            o.convert(b"x", ".pptx", "pdf")
        assert time.monotonic() - t0 < 0.1
        assert (e.value.status, e.value.code) == (429, "BUSY")
        assert helper_ok == []
    finally:
        o._lock.release()


def test_helper_timeout_stops_office_and_504(office, monkeypatch):
    """헬퍼 타임아웃 → 504 + soffice 를 내리기만 한다(재기동 대기로 응답이 늘어지지 않게). 다음 변환이 다시 띄운다."""
    o, events = office

    def run(cmd, timeout, cwd, on_fail=None):
        raise ConvertError(504, "TIMEOUT")

    monkeypatch.setattr(proc, "run", run)
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".pptx", "pdf")
    assert (e.value.status, e.value.code) == (504, "TIMEOUT")
    assert events == ["stop"]
    # 실패 후에도 슬롯이 풀려 다음 요청을 받을 수 있어야 한다
    assert o._lock.acquire(blocking=False)
    o._lock.release()


def _helper_exits(monkeypatch, rc: int) -> None:
    """proc.run 대체 — 헬퍼가 rc 로 끝난 것처럼 on_fail 매핑을 그대로 태운다."""
    def run(cmd, timeout, cwd, on_fail=None):
        raise on_fail(rc)

    monkeypatch.setattr(proc, "run", run)


def test_helper_bad_input_is_422_without_restart(office, monkeypatch):
    """헬퍼 종료코드 2(문서 로드 실패) — 입력 문제이므로 422, 멀쩡한 soffice 는 건드리지 않는다."""
    o, events = office
    _helper_exits(monkeypatch, 2)
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".doc", "pdf")
    assert (e.value.status, e.value.code) == (422, "CONVERT_FAILED")
    assert events == []


@pytest.mark.parametrize("rc", [3, 1, -9])
def test_helper_infra_failure_is_503_and_stops(office, monkeypatch, rc):
    """연결 실패(3)·import 실패(1)·시그널 사망(음수) — 인프라 오류 503, soffice 를 내린다."""
    o, events = office
    _helper_exits(monkeypatch, rc)
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".pptx", "pdf")
    assert (e.value.status, e.value.code) == (503, "UNAVAILABLE")
    assert events == ["stop"]


def test_bad_input_but_soffice_died_is_503(office, monkeypatch):
    """rc 2 라도 그 사이 soffice 가 죽었으면(변환 중 크래시) 인프라 오류 503 + 내리기."""
    o, events = office
    _helper_exits(monkeypatch, 2)
    monkeypatch.setattr(o, "_alive", lambda: False)
    o._proc = RUNNING
    # 지연 기동 판정도 _alive 를 보므로 기동은 일어나고(start), 헬퍼 실패 후 죽음 확인 → stop
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".pptx", "pdf")
    assert (e.value.status, e.value.code) == (503, "UNAVAILABLE")
    assert events == ["start", "stop"]


def test_lazy_start_failure_is_503(office, helper_ok, monkeypatch):
    """꺼진 soffice 를 띄우려다 실패하면 503 UNAVAILABLE, 헬퍼는 부르지 않는다."""
    o, _ = office
    o._proc = None

    def broken_start(wait=soffice.START_WAIT_SEC):
        raise ConvertError(503, "UNAVAILABLE")

    monkeypatch.setattr(o, "start", broken_start)
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".pptx", "pdf")
    assert (e.value.status, e.value.code) == (503, "UNAVAILABLE")
    assert helper_ok == []


def test_shared_deadline_covers_start_and_helper(office, monkeypatch):
    """지연 기동이 쓴 시간만큼 헬퍼 타임아웃이 줄어 요청 전체가 TIMEOUT_SEC 안에 묶인다."""
    o, _ = office
    o._proc = None
    monkeypatch.setattr(config, "TIMEOUT_SEC", 2)
    seen = {}

    def slow_start(wait=soffice.START_WAIT_SEC):
        seen["wait"] = wait
        time.sleep(0.3)
        o._proc = RUNNING

    def run(cmd, timeout, cwd, on_fail=None):
        seen["timeout"] = timeout
        Path(cmd[cmd.index("--out") + 1]).write_bytes(b"ok")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    monkeypatch.setattr(o, "start", slow_start)
    monkeypatch.setattr(proc, "run", run)
    o.convert(b"x", ".pptx", "pdf")
    assert seen["wait"] <= 2
    assert seen["timeout"] <= 1.75


def test_start_consuming_budget_is_503_without_helper(office, helper_ok, monkeypatch):
    """기동이 예산을 거의 다 쓰면 헬퍼를 띄우지 않고 503 — 뻔한 타임아웃으로 멀쩡한 soffice 를 죽이지 않는다."""
    o, events = office
    o._proc = None
    monkeypatch.setattr(config, "TIMEOUT_SEC", 1)

    def slow_start(wait=soffice.START_WAIT_SEC):
        events.append("start")
        time.sleep(0.3)
        o._proc = RUNNING

    monkeypatch.setattr(o, "start", slow_start)
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".pptx", "pdf")
    assert (e.value.status, e.value.code) == (503, "UNAVAILABLE")
    assert helper_ok == []
    assert events == ["start"]


def test_health_busy_is_ok_without_start(office):
    """변환 중(락 점유)이면 probe 를 막지 않고 살아 있음으로 본다."""
    o, events = office
    o._proc = None
    assert o._lock.acquire(blocking=False)
    try:
        assert o.health() is True
    finally:
        o._lock.release()
    assert events == []


def test_health_restarts_dead_office(office):
    o, events = office
    o._proc = None
    assert o.health() is True
    assert events == ["start"]


def test_health_false_when_restart_fails(office, monkeypatch):
    o, _ = office
    o._proc = None

    def broken_start(wait=soffice.START_WAIT_SEC):
        raise ConvertError(503, "UNAVAILABLE")

    monkeypatch.setattr(o, "start", broken_start)
    assert o.health() is False


def test_missing_soffice_binary_convert_is_503(missing_binary, helper_ok):
    """실제 Office — soffice 실행 파일이 없으면 변환은 503 UNAVAILABLE(손상 파일 422 와 구분)."""
    o = missing_binary
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".pptx", "pdf")
    assert (e.value.status, e.value.code) == (503, "UNAVAILABLE")
    assert helper_ok == []
    assert o.health() is False


def test_restart_after_n_conversions(office, helper_ok, monkeypatch):
    o, events = office
    monkeypatch.setattr(config, "SOFFICE_RESTART_EVERY", 2)
    o.convert(b"a", ".pptx", "pdf")
    o.convert(b"b", ".pptx", "pdf")
    assert events == []
    o.convert(b"c", ".pptx", "pdf")
    # 3번째 변환의 헬퍼 호출 전에 stop → start
    assert events == ["stop", "start"]
    assert len(helper_ok) == 3


def test_lazy_start_when_not_running(office, helper_ok):
    """기동 실패(개발 호스트 등)로 꺼져 있으면 첫 변환 때 기동한다."""
    o, events = office
    o._proc = None
    o.convert(b"x", ".pptx", "pdf")
    assert events == ["start"]


def test_helper_cmd_uses_system_python(office, helper_ok, monkeypatch):
    o, _ = office
    monkeypatch.setattr(config, "UNO_PYTHON", "/usr/bin/python3")
    out = o.convert(b"DOC-BYTES", ".pptx", "txt")
    assert out == b"%PDF-fake"
    cmd = helper_ok[0]
    assert cmd[0] == config.UNO_PYTHON
    helper = Path(cmd[1])
    assert helper.is_absolute() and helper.name == "uno_helper.py" and helper.is_file()
    assert cmd[2:4] == ["--port", str(config.SOFFICE_PORT)]
    assert cmd[cmd.index("--target") + 1] == "txt"
    assert Path(cmd[cmd.index("--in") + 1]).suffix == ".pptx"


@pytest.mark.parametrize("filter_name", ["Impress MS PowerPoint 2007 XML", None])
def test_helper_import_filter_flag(office, helper_ok, filter_name):
    """filter_name 을 주면 헬퍼에 --filter 로 넘겨 그 필터로만 열고, 안 주면 --filter 없이 호출(자동 판별)."""
    o, _ = office
    o.convert(b"x", ".pptx", "pdf", filter_name=filter_name)
    cmd = helper_ok[0]
    if filter_name:
        assert cmd[cmd.index("--filter") + 1] == filter_name
    else:
        assert "--filter" not in cmd


def test_helper_receives_input_bytes_and_cleans_tmp(office, monkeypatch, tmp_path):
    """헬퍼가 받는 입력 파일 = 요청 바이트, 변환 후 요청 임시 디렉터리는 지워진다."""
    o, _ = office
    seen = {}

    def run(cmd, timeout, cwd, on_fail=None):
        src = Path(cmd[cmd.index("--in") + 1])
        seen["data"] = src.read_bytes()
        seen["dir"] = src.parent
        Path(cmd[cmd.index("--out") + 1]).write_bytes(b"ok")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    monkeypatch.setattr(proc, "run", run)
    o.convert(b"\x00ORIG", ".doc", "pdf")
    assert seen["data"] == b"\x00ORIG"
    assert not seen["dir"].exists()


def test_missing_output_is_convert_failed(office, monkeypatch):
    """헬퍼가 0 으로 끝났는데 결과 파일이 없으면 422."""
    o, _ = office
    monkeypatch.setattr(proc, "run",
                        lambda cmd, timeout, cwd, on_fail=None: subprocess.CompletedProcess(cmd, 0, b"", b""))
    with pytest.raises(ConvertError) as e:
        o.convert(b"x", ".pptx", "pdf")
    assert (e.value.status, e.value.code) == (422, "CONVERT_FAILED")


def test_start_with_missing_binary_is_unavailable(missing_binary):
    """soffice 실행 파일이 없으면 기동은 503 UNAVAILABLE 로 실패한다(앱 기동은 이를 로그만 남긴다)."""
    with pytest.raises(ConvertError) as e:
        missing_binary.start()
    assert (e.value.status, e.value.code) == (503, "UNAVAILABLE")
    assert not missing_binary._running


class _DeadWrapper:
    """이미 끝난 soffice 래퍼 흉내 — stop() 이 kill·wait 를 거쳐 포트 대기까지 가게 한다."""
    pid = 999_999

    def poll(self):
        return 0

    def wait(self, timeout=None):
        return 0


def test_stop_waits_until_port_closed(monkeypatch):
    """래퍼가 끝나도 soffice.bin 이 잠깐 포트를 쥐고 있을 수 있다 — stop() 은 포트가 닫힌 뒤에야 돌아온다.

    안 기다리면 바로 이은 start() 가 죽어 가는 이전 인스턴스의 포트를 보고 기동 완료로 착각한다(엔진 테스트 간헐 실패 원인).
    """
    polls = iter([True, True, True, False])
    seen: list[bool] = []
    monkeypatch.setattr(soffice, "_port_open", lambda port: seen.append(next(polls)) or seen[-1])
    killed: list[int] = []
    monkeypatch.setattr(soffice.os, "killpg", lambda pgid, sig: killed.append(pgid))
    o = soffice.Office()
    o._proc = _DeadWrapper()
    o.stop()
    assert seen == [True, True, True, False]
    # 래퍼가 끝났어도 그룹에 남은 soffice.bin 을 위해 SIGKILL 은 보낸다
    assert killed == [_DeadWrapper.pid]
    assert o._proc is None


def test_start_refuses_when_port_already_open(monkeypatch, tmp_path):
    """띄우기 전에 포트가 이미 열려 있으면(남의 소켓) 기동 판정을 믿을 수 없으므로 띄우지 않고 503."""
    monkeypatch.setenv("TMPDIR", str(tmp_path))
    spawned: list[list[str]] = []
    monkeypatch.setattr(soffice.subprocess, "Popen", lambda cmd, **kw: spawned.append(cmd))
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        s.listen()
        monkeypatch.setattr(config, "SOFFICE_PORT", s.getsockname()[1])
        with pytest.raises(ConvertError) as e:
            soffice.Office().start()
    assert (e.value.status, e.value.code) == (503, "UNAVAILABLE")
    assert spawned == []


@pytest.mark.engine
def test_real_stop_releases_port_before_return(real_office):
    """실물: stop() 이 돌아온 시점엔 UNO 포트가 이미 닫혀 있다(다음 start 가 이전 인스턴스에 속지 않게)."""
    real_office.start()
    real_office.convert(_one_slide_pptx(), ".pptx", "pdf")
    real_office.stop()
    assert not soffice._port_open(config.SOFFICE_PORT)


@pytest.mark.engine
def test_uno_helper_disables_char_distance():
    """한글↔영문·숫자 혼합 문장에 LibreOffice 자동 공백이 끼지 않아야 한다."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text_frame.text = "AI 시대, 2,603명"
    buf = io.BytesIO()
    prs.save(buf)

    o = soffice.Office()
    try:
        o.start()
        pdf = o.convert(buf.getvalue(), ".pptx", "pdf")
    finally:
        o.stop()
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        text = "".join(page.get_text() for page in doc)
    assert "2,603명" in text
    assert "시대," in text
    assert "2,603 명" not in text


# ---- 실제 soffice 수명(start/stop/재기동) — 단위 테스트는 start/stop 을 가짜로 바꾸므로 여기서 실물로 검증 ----

def _one_slide_pptx(text: str = "재기동 테스트") -> bytes:
    """텍스트 상자 하나짜리 1장 PPTX."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _session_pids(sid: int) -> list[int]:
    """/proc 를 훑어 세션 sid 에 속한 살아 있는 프로세스 pid 목록(좀비 제외). slim 이미지엔 pgrep 이 없다."""
    pids = []
    for p in Path("/proc").iterdir():
        if not p.name.isdigit():
            continue
        try:
            stat = (p / "stat").read_text()
        except OSError:
            continue  # 그 사이 사라진 프로세스
        # stat 형식: pid (comm) state ppid pgrp session ... — comm 에 공백이 있을 수 있어 마지막 ')' 뒤를 자른다
        fields = stat[stat.rindex(")") + 2:].split()
        if fields[0] != "Z" and int(fields[3]) == sid:
            pids.append(int(p.name))
    return pids


def _soffice_bin_pid(sid: int, timeout: float = 15) -> int:
    """세션 sid 안의 실제 오피스 프로세스(soffice.bin) pid. 래퍼(oosplash)가 아닌 UNO 소켓을 연 프로세스.

    새 프로필로 처음 뜰 때 soffice.bin 이 프로필을 만들고 스스로 한 번 재시작(oosplash 가 다시 띄움)하므로
    그 사이 잠깐 soffice.bin 이 없을 수 있다 — 나타날 때까지 기다린다.
    """
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        for pid in _session_pids(sid):
            try:
                cmd = Path(f"/proc/{pid}/cmdline").read_bytes()
            except OSError:
                continue
            if b"soffice.bin" in cmd:
                return pid
        time.sleep(0.1)
    raise AssertionError(f"세션 {sid} 에 soffice.bin 이 없음")


def _wait_gone(pid: int, timeout: float = 10) -> None:
    """pid 가 사라지거나 좀비가 될 때까지 기다린다."""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            if Path(f"/proc/{pid}/stat").read_text().split(") ")[1].startswith("Z"):
                return
        except (OSError, IndexError):
            return
        time.sleep(0.1)
    raise AssertionError(f"pid {pid} 가 {timeout}s 안에 끝나지 않음")


@pytest.fixture
def real_office(monkeypatch, tmp_path):
    """실제 soffice 를 쓰는 새 Office — TMPDIR 을 테스트 전용으로 두고, 끝나면 반드시 내린다."""
    monkeypatch.setenv("TMPDIR", str(tmp_path))
    o = soffice.Office()
    yield o
    o.stop()


@pytest.mark.engine
def test_real_start_convert_stop_leaves_no_process(real_office):
    """start → 변환 성공 → stop 하면 soffice 세션의 프로세스가 하나도 남지 않는다."""
    real_office.start()
    sid = real_office._proc.pid
    assert _session_pids(sid), "기동 직후 soffice 세션이 있어야 함"
    pdf = real_office.convert(_one_slide_pptx(), ".pptx", "pdf")
    assert pdf.startswith(b"%PDF")

    real_office.stop()
    assert real_office._proc is None and not real_office._running
    deadline = time.monotonic() + 5
    while _session_pids(sid) and time.monotonic() < deadline:
        time.sleep(0.1)
    assert _session_pids(sid) == []


@pytest.mark.engine
def test_real_soffice_killed_group_restarts_on_next_convert(real_office):
    """soffice 프로세스 그룹이 통째로 죽으면(크래시·OOM kill) 다음 변환이 새로 띄워 성공한다."""
    import os
    import signal

    real_office.start()
    old = real_office._proc.pid
    assert real_office.convert(_one_slide_pptx("첫 변환"), ".pptx", "pdf").startswith(b"%PDF")

    os.killpg(old, signal.SIGKILL)
    _wait_gone(old)

    pdf = real_office.convert(_one_slide_pptx("재기동후"), ".pptx", "pdf")
    assert pdf.startswith(b"%PDF")
    assert real_office._proc.pid != old
    # 공백 없는 문구로 확인한다 — 한글 사이 공백은 라틴 글꼴 span 으로 따로 그려져 pymupdf 추출에서 빠진다(렌더링은 정상)
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        assert "재기동후" in doc[0].get_text()


@pytest.mark.engine
def test_real_soffice_bin_killed_recovers_within_one_request(real_office):
    """래퍼는 살아 있고 soffice.bin 만 죽은 경우 — 다음 요청은 (포트 확인으로 재기동해) 성공하거나 503 후 내려지고,
    그다음 요청은 반드시 성공."""
    import os
    import signal

    real_office.start()
    sid = real_office._proc.pid
    os.kill(_soffice_bin_pid(sid), signal.SIGKILL)

    try:
        first = real_office.convert(_one_slide_pptx(), ".pptx", "pdf")
        assert first.startswith(b"%PDF")
    except ConvertError as e:
        assert (e.status, e.code) == (503, "UNAVAILABLE")
    assert real_office.convert(_one_slide_pptx(), ".pptx", "pdf").startswith(b"%PDF")


@pytest.mark.engine
def test_real_bad_input_is_422_and_soffice_not_restarted(real_office):
    """확장자만 바꾼 파일(평문을 .pptx 로) → 422, 상주 soffice 는 같은 PID 로 살아 있고 다음 변환도 그대로 성공."""
    real_office.start()
    proc_before = real_office._proc
    pid = proc_before.pid
    with pytest.raises(ConvertError) as e:
        real_office.convert("그냥 텍스트\nplain\n".encode("utf-8"), ".pptx", "pdf",
                            filter_name="Impress MS PowerPoint 2007 XML")
    assert (e.value.status, e.value.code) == (422, "CONVERT_FAILED")
    assert real_office._proc is proc_before and real_office._proc.pid == pid
    assert real_office._proc.poll() is None
    assert real_office.convert(_one_slide_pptx(), ".pptx", "pdf").startswith(b"%PDF")
    assert real_office._proc.pid == pid


@pytest.mark.engine
def test_real_helper_connect_failure_exit_3(tmp_path):
    """soffice 가 없는 포트로 붙으면 헬퍼는 연결 실패 종료코드 3(입력 문제 2 와 구분)."""
    src = tmp_path / "in.pptx"
    src.write_bytes(_one_slide_pptx())
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    cp = subprocess.run([config.UNO_PYTHON, str(soffice.HELPER), "--port", str(port), "--in", str(src),
                         "--out", str(tmp_path / "out.pdf"), "--target", "pdf"],
                        capture_output=True, timeout=60)
    assert cp.returncode == 3, cp.stderr


@pytest.mark.engine
def test_real_calc_text_bulk_read_values(real_office):
    """시트 텍스트 일괄 읽기 — 빈 셀은 빈 칸, 정수는 '.0' 없이, 수식 오류 셀은 'None' 이 아닌 표시 문자열."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["C1"], ws["D1"], ws["A2"] = "이름", 3, 2.5, "=1/0"
    buf = io.BytesIO()
    wb.save(buf)
    text = real_office.convert(buf.getvalue(), ".xlsx", "txt").decode("utf-8")
    lines = text.split("\n")
    assert lines[0] == "이름\t\t3\t2.5"
    assert "None" not in text
    assert lines[1].startswith("#")
