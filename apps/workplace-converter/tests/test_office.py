"""Office 엔진 테스트 — PPTX·PPT·DOC → PDF, PPTX·PPT·DOC·XLS → 텍스트.

단위 테스트는 OFFICE.convert 를 가짜로 바꿔 mime↔확장자·가져오기 필터 매핑, 쪽수 계산, 오류 매핑만 검증한다.
`@pytest.mark.engine` 테스트는 soffice·python3-uno 가 있는 컨테이너 test 타깃에서만 돈다.
"""
import io
import os
from pathlib import Path

import pymupdf
import pytest

from app import main
from app.engines import office
from app.errors import ConvertError

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PPT = "application/vnd.ms-powerpoint"
DOC = "application/msword"
XLS = "application/vnd.ms-excel"
FIXTURES = Path(__file__).parent / "fixtures"


def _tiny_pdf(pages: int) -> bytes:
    """쪽수가 정해진 최소 PDF 바이트 — 가짜 soffice 출력용."""
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"page {i + 1}")
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def fake_office(monkeypatch):
    """OFFICE.convert 대체 — 호출 인자를 기록하고 state['out'] 바이트를 돌려준다."""
    state = {"calls": [], "out": _tiny_pdf(1)}

    def convert(data, suffix, target, filter_name=None):
        state["calls"].append({"data": data, "suffix": suffix, "target": target, "filter": filter_name})
        return state["out"]

    monkeypatch.setattr(office.OFFICE, "convert", convert)
    return state


# ---- 단위 ----

@pytest.mark.parametrize("mime,suffix", [(PPTX, ".pptx"), (PPT, ".ppt"), (DOC, ".doc")])
def test_office_suffix_by_mime(fake_office, mime, suffix):
    office.to_pdf(b"SRC", mime)
    call = fake_office["calls"][0]
    assert (call["data"], call["suffix"], call["target"]) == (b"SRC", suffix, "pdf")


@pytest.mark.parametrize("mime,filter_name", [
    (PPTX, "Impress MS PowerPoint 2007 XML"),
    (PPT, "MS PowerPoint 97"),
    (DOC, "MS Word 97"),
    (XLS, "MS Excel 97"),
])
def test_office_import_filter_by_mime(fake_office, mime, filter_name):
    """확장자만 바꾼 파일이 자동 판별로 통과하지 않게 mime 별 가져오기 필터를 명시한다."""
    fake_office["out"] = b"text"
    office.to_text(b"SRC", mime)
    call = fake_office["calls"][0]
    assert call["filter"] == filter_name
    assert call["suffix"] == office._SPEC[mime][0]
    assert call["target"] == "txt"


def test_office_pdf_page_count_from_pymupdf(fake_office):
    fake_office["out"] = _tiny_pdf(2)
    result = office.to_pdf(b"x", PPTX)
    assert result.page_count == 2
    assert result.data == fake_office["out"]


def test_office_corrupt_pdf_output_is_422(fake_office):
    """soffice 가 PDF 가 아닌 바이트를 내면 500 이 아니라 422 CONVERT_FAILED."""
    fake_office["out"] = b"not a pdf"
    with pytest.raises(ConvertError) as e:
        office.to_pdf(b"x", DOC)
    assert (e.value.status, e.value.code) == (422, "CONVERT_FAILED")


# 쪽 없는 최소 PDF — pymupdf 는 0쪽 문서 저장을 거부하므로 손으로 쓴다
_EMPTY_PDF = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
              b"2 0 obj<</Type/Pages/Kids[]/Count 0>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n")


def test_office_zero_page_pdf_is_422(fake_office):
    """soffice 가 0쪽 PDF 를 내면 성공이 아니라 422 CONVERT_FAILED."""
    fake_office["out"] = _EMPTY_PDF
    with pytest.raises(ConvertError) as e:
        office.to_pdf(b"x", PPTX)
    assert (e.value.status, e.value.code) == (422, "CONVERT_FAILED")


@pytest.mark.parametrize("bom", [b"", b"\xef\xbb\xbf"], ids=["plain", "bom"])
def test_office_text_is_utf8_decoded(fake_office, bom):
    """UTF-8 디코드 — Writer 텍스트 필터가 붙이는 BOM 은 떼어낸다."""
    fake_office["out"] = bom + "AI 시대, 2,603명".encode("utf-8")
    assert office.to_text(b"x", XLS) == "AI 시대, 2,603명"


def test_office_engines_registered():
    """main 레지스트리: PDF 는 PPTX·PPT·DOC, 텍스트는 거기에 XLS 추가."""
    for mime in (PPTX, PPT, DOC):
        assert main.PDF_ENGINES[mime] is office.to_pdf
        assert main.TEXT_ENGINES[mime] is office.to_text
    assert XLS not in main.PDF_ENGINES
    assert main.TEXT_ENGINES[XLS] is office.to_text


# ---- 엔진(컨테이너) ----

@pytest.fixture
def live_office():
    """실제 상주 soffice — 테스트 끝나면 내린다."""
    yield office.OFFICE
    office.OFFICE.stop()


def _pptx(slides: int) -> bytes:
    """python-pptx 로 슬라이드 수가 정해진 PPTX 를 만든다."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text_frame.text = f"슬라이드 {i + 1}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _fixture(name: str) -> bytes:
    """Task 5 의 이미지 안에서 생성되는 픽스처 — 없으면 건너뛴다."""
    path = FIXTURES / name
    if not path.is_file():
        pytest.skip(f"픽스처 {name} 없음 — 컨테이너 이미지 빌드 시 생성됨")
    return path.read_bytes()


@pytest.mark.engine
def test_pptx_to_pdf_pages_match_slides(live_office):
    result = office.to_pdf(_pptx(3), PPTX)
    assert result.page_count == 3
    assert result.data.startswith(b"%PDF")


@pytest.mark.engine
def test_doc_to_pdf(live_office):
    result = office.to_pdf(_fixture("sample.doc"), DOC)
    assert result.page_count >= 1
    assert result.data.startswith(b"%PDF")


@pytest.mark.engine
def test_xls_to_text_contains_cell(live_office):
    text = office.to_text(_fixture("sample.xls"), XLS)
    assert "변환 테스트" in text


@pytest.mark.engine
@pytest.mark.parametrize("bad,mime", [
    (os.urandom(4096), DOC),  # 손상 파일
    ("그냥 텍스트 파일입니다\nplain text\n".encode("utf-8"), PPTX),  # 확장자만 바꾼 평문 — 자동 판별로 통과하면 안 됨
], ids=["corrupt-doc", "text-renamed-pptx"])
def test_bad_input_is_422_and_next_ok(live_office, bad, mime):
    """입력 문제는 422, 직후 정상 PPTX 는 그대로 성공(soffice 가 멀쩡히 남는다)."""
    with pytest.raises(ConvertError) as e:
        office.to_pdf(bad, mime)
    assert (e.value.status, e.value.code) == (422, "CONVERT_FAILED")
    assert office.to_pdf(_pptx(2), PPTX).page_count == 2


@pytest.mark.engine
def test_doc_to_text_without_bom(live_office):
    """Writer 텍스트 필터 경로 — 본문이 나오고 BOM 이 앞에 남지 않는다."""
    text = office.to_text(_fixture("sample.doc"), DOC)
    assert not text.startswith("\ufeff")
    assert "2,603명" in text and "lazy dog" in text


@pytest.mark.engine
def test_pptx_to_text_joins_slides(live_office):
    """Impress 도형 텍스트 경로 — 슬라이드 순서대로 빈 줄로 잇는다."""
    assert office.to_text(_pptx(2), PPTX) == "슬라이드 1\n\n슬라이드 2"
